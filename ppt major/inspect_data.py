import sys
import pptx
from pptx import Presentation
import docx

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'c:\Users\harsh\OneDrive\Documents\fy_mj_project\Project Idea Presentaion Template.pptx')
print("=== Project Idea Presentaion Template.pptx Details ===")
print("Slide width:", prs.slide_width, "height:", prs.slide_height)

for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        print(f" Shape {s_idx}: name='{shape.name}', type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                print(f"   P{p_idx}: '{p.text}'")
        if shape.has_table:
            table = shape.table
            print(f"   Table: {len(table.rows)} rows x {len(table.columns)} cols")
            for r in table.rows:
                row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
                print(f"     Row: {row_txt}")

print("\n" + "="*50 + "\n")

prs2 = Presentation(r'c:\Users\harsh\OneDrive\Documents\fy_mj_project\Agentic_AI_Cybercrime_Investigation_Synopsis.pptx')
print("=== Agentic_AI_Cybercrime_Investigation_Synopsis.pptx Details ===")
for i, slide in enumerate(prs2.slides):
    print(f"\n--- Slide {i+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            text = shape.text_frame.text.strip().replace('\n', ' ')
            if text:
                print(f"   Shape {s_idx} ({shape.name}): {text[:150]}")

print("\n" + "="*50 + "\n")
print("=== Checking DOCX files ===")
doc_path = r'c:\Users\harsh\OneDrive\Documents\fy_mj_project\Agentic_AI_Cybercrime_Investigation_Synopsis.docx'
doc = docx.Document(doc_path)
for p in doc.paragraphs:
    if p.text.strip():
        print("DOCX P:", p.text[:120])
