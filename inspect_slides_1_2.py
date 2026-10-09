import sys
import pptx
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'c:\Users\harsh\OneDrive\Documents\fy_mj_project\Project Idea Presentaion Template.pptx')

for i in [0, 1]:
    slide = prs.slides[i]
    print(f"=== SLIDE {i+1} ===")
    for j, shape in enumerate(slide.shapes):
        print(f"Shape [{j}] '{shape.name}' type={shape.shape_type} left={shape.left} top={shape.top} width={shape.width} height={shape.height}")
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                print(f"  Text: '{p.text}'")
        if shape.has_table:
            t = shape.table
            print(f"  Table {len(t.rows)}x{len(t.columns)}")
            for r_idx, row in enumerate(t.rows):
                for c_idx, cell in enumerate(row.cells):
                    print(f"    Cell ({r_idx},{c_idx}): '{cell.text.strip().replace('\n', ' ')}'")
