import sys
import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation(r'c:\Users\harsh\OneDrive\Documents\fy_mj_project\Project Idea Presentaion Template.pptx')
print(f"Slide dimensions: {prs.slide_width} x {prs.slide_height} (Emu)")
print(f"Width in inches: {prs.slide_width / 914400:.2f}, Height in inches: {prs.slide_height / 914400:.2f}")

os.makedirs('extracted_template_assets', exist_ok=True)

for i, slide in enumerate(prs.slides):
    print(f"\n================ SLIDE {i+1} ================")
    if slide.has_notes_slide:
        print(f"Notes: {slide.notes_slide.notes_text_frame.text}")
    for j, shape in enumerate(slide.shapes):
        print(f"\n  Shape [{j}] Name: '{shape.name}', Type: {shape.shape_type}")
        print(f"  Pos: Left={shape.left}, Top={shape.top}, Width={shape.width}, Height={shape.height}")
        
        # Save image shapes if any
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = shape.image
            img_filename = f"extracted_template_assets/slide_{i+1}_img_{j}.{image.ext}"
            with open(img_filename, 'wb') as f:
                f.write(image.blob)
            print(f"  [Saved Picture to {img_filename}]")

        if shape.has_text_frame:
            tf = shape.text_frame
            for p_idx, p in enumerate(tf.paragraphs):
                p_text = p.text
                if p_text.strip():
                    runs_info = []
                    for r in p.runs:
                        font_color = None
                        if r.font.color and r.font.color.type == 1:
                            font_color = str(r.font.color.rgb)
                        runs_info.append(f"(text='{r.text}', font='{r.font.name}', sz={r.font.size}, bold={r.font.bold}, color={font_color})")
                    print(f"    P{p_idx} Align={p.alignment}: '{p_text}' | Runs: {runs_info}")
        
        if shape.has_table:
            table = shape.table
            print(f"    Table ({len(table.rows)}x{len(table.columns)}):")
            for r_idx, row in enumerate(table.rows):
                row_str = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(f"      Row {r_idx}: {row_str}")
